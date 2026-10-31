"""Run: .venv/bin/python -m unittest discover -s tests -v

The intake (numbers -> values), the shared options and the workbook writer. Everything is written to a temp dir."""

import json
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
import intake as IN  # noqa: E402
import options as OPT  # noqa: E402
import write_questionnaire as WQ  # noqa: E402
from paths import AI_AREAS, AI_SCOPE_RULES  # noqa: E402


def args(**over):
    base = dict(replace=False, org="Acme Corp", domain="1", certs="2 3", ai_in_use="2", role="1", built_on_model=None,
                given_to_others=None, ai_usage="1 3", ai_components="2", csf_assessment_done="1", iso42001_held="2",
                date="2026-10-14", runs_dir=None)
    base.update(over)
    return type("A", (), base)()


class Options(unittest.TestCase):
    def test_every_ai_area_is_reachable(self):
        # every AI area comes from a menu or ai_area_rules; automatic ones (always, any_ai) are never also on a menu.
        # A with-rule may add a menu item too (agents add MCP / tools, which can also be ticked).
        areas = {a["id"] for a in json.loads(AI_AREAS.read_text())["ai_areas"]}
        menu = {q["key"]: [o["value"] for o in q["options"]] for q in OPT.INTAKE_QUESTIONS}
        r = json.loads(AI_SCOPE_RULES.read_text())["ai_area_rules"]
        menus, auto = set(menu["ai_usage"]) | set(menu["ai_components"]), set(r["always"]) | set(r["any_ai"])
        added = {a for adds in list(r["with"].values()) + list(r["role"].values()) for a in adds}
        self.assertEqual(menus | auto | added, areas)
        self.assertFalse(menus & auto)
        self.assertTrue(set(r["role"]) <= set(menu["role"]))
        self.assertTrue(set(r["with"]) <= areas)

    def test_client_roles_cover_every_answer_and_only_known_roles(self):
        rules, qrules = json.loads(AI_SCOPE_RULES.read_text()), json.loads(OPT.AI_QUESTION_RULES.read_text())
        menu = {q["key"]: [o["value"] for o in q["options"]] for q in OPT.INTAKE_QUESTIONS}
        roles = {k: v for k, v in rules["client_roles"].items() if k != "description"}
        self.assertEqual(set(roles), set(menu["role"]) | {"none", "unsure"})
        self.assertTrue({r for v in roles.values() for r in v} <= set(qrules["roles"]))

    def test_every_domain_has_a_weight_key_and_every_cert_a_file(self):
        self.assertEqual(set(OPT.WEIGHT_KEY), set(OPT.DOMAINS))
        self.assertEqual(set(OPT.CERT_FILES), set(OPT.CERTS))


class Intake(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def go(self, **over):
        a = args(runs_dir=str(self.tmp), **over)
        try:
            IN.write(a)
            return 0
        except SystemExit as e:
            return e.code

    def run_dir(self):
        return self.tmp / "Acme_Corp_2026-10-14"

    def usage(self):
        return json.loads((self.run_dir() / "ai_usage.json").read_text())

    def test_numbers_become_ai_areas_and_the_iso_flag_is_derived(self):
        self.assertEqual(self.go(), 0)
        intake = json.loads((self.run_dir() / "intake.json").read_text())
        self.assertEqual((intake["domain"], intake["certs"]), ("finance", ["iso27001", "pci_dss"]))
        usage = self.usage()
        self.assertEqual(usage["ai_areas"], ["internal_process", "customer_facing", "invisible_usage", "input", "output",
                                             "mcp_tools"])
        self.assertEqual((usage["role"], usage["producer"]), ("consumer", None))
        self.assertFalse({"rung", "usage_contexts", "modifiers", "hosts_model"} & set(usage))   # the old fields are gone
        self.assertEqual(usage["engagement_flags"], {"csf_assessment_done": True, "iso27001_held": True,
                                                     "iso42001_held": False})

    def test_no_ai_keeps_only_invisible_usage_and_pci_only_leaves_iso_false(self):
        self.assertEqual(self.go(ai_in_use="1", role=None, ai_usage=None, ai_components=None, certs="3"), 0)
        usage = self.usage()
        self.assertEqual(usage["ai_areas"], ["invisible_usage"])
        self.assertEqual((usage["role"], usage["producer"]), (None, None))
        self.assertFalse(usage["engagement_flags"]["iso27001_held"])

    def test_producer_and_agents_bring_their_areas(self):
        # build, built on an existing model, not given to others; agents and a hosted model
        self.assertEqual(self.go(role="2", built_on_model="1", given_to_others="2", ai_usage="", ai_components="3 4"), 0)
        usage = self.usage()
        for area in ("model_building", "model", "test_tuning_data", "agent", "mcp_tools", "hosting"):
            self.assertIn(area, usage["ai_areas"])
        self.assertEqual(usage["producer"], {"built_on_existing_model": True, "given_to_others": False})
        self.assertEqual(usage["role"], "producer")

    def test_both_keeps_the_usage_areas_and_adds_the_producer_ones(self):
        self.assertEqual(self.go(role="3", built_on_model="2", given_to_others="1"), 0)
        usage = self.usage()
        self.assertEqual(usage["role"], "both")
        self.assertTrue({"internal_process", "customer_facing", "model_building", "model"} <= set(usage["ai_areas"]))
        self.assertEqual(usage["producer"], {"built_on_existing_model": False, "given_to_others": True})

    def test_blank_menus_are_allowed(self):
        self.assertEqual(self.go(ai_usage="", ai_components=""), 0)
        usage = self.usage()
        self.assertEqual(usage["ai_areas"], ["invisible_usage", "input", "output"])
        self.assertEqual(usage["role"], "consumer")

    def test_bad_input_is_rejected_and_writes_nothing(self):
        for bad in (dict(domain="9"), dict(domain="x"), dict(domain="1 2"), dict(certs="1 2"), dict(certs=""),
                    dict(ai_usage="9"), dict(ai_components="1 7"), dict(role=None), dict(role="4"),
                    dict(built_on_model="1"),                                 # producer follow-up for a consumer
                    dict(role="2", built_on_model=None, given_to_others="2"),  # follow-up missing for a producer
                    dict(ai_in_use="1"),                                      # questions answered without "some"
                    dict(ai_in_use="3", role=None, ai_usage=None, ai_components="1"),
                    dict(org="!!!")):
            self.assertEqual(self.go(**bad), 1, bad)
            self.assertFalse(self.run_dir().exists(), bad)

    def test_an_existing_intake_is_not_overwritten_without_replace(self):
        self.assertEqual(self.go(), 0)
        self.assertEqual(self.go(domain="2"), 1)
        self.assertEqual(json.loads((self.run_dir() / "intake.json").read_text())["domain"], "finance")
        self.assertEqual(self.go(domain="2", replace=True), 0)
        self.assertEqual(json.loads((self.run_dir() / "intake.json").read_text())["domain"], "healthcare")

    def test_downstream_scripts_find_domain_and_certs(self):
        self.go()
        self.assertEqual(OPT.load_intake(self.run_dir())["certs"], ["iso27001", "pci_dss"])
        with self.assertRaises(SystemExit):
            OPT.load_intake(self.tmp / "missing")


class RecordLog(unittest.TestCase):
    """record_judgements.py log (assessor patch merge) now goes through audit.record_agent_output."""

    def test_a_bad_patch_is_rejected_and_nothing_is_written(self):
        import record_judgements as RJ
        run = Path(tempfile.mkdtemp())
        entry = {"subcategory_id": "GV.RM-01", "function": "GV", "outcome": "o", "question_type": "full_discovery",
                 "questions": [{"question_text": "How is risk managed?"}], "evidence_prompt": "", "scope_reason": "",
                 "confidence_flag": "", "assessor_note": None}
        (run / "draft.json").write_text(json.dumps([entry]))
        (run / "patch.json").write_text(json.dumps([entry]))          # identical to the draft, no note: invalid
        argv = [str(run), "--step", "assessor", "--base", "draft.json", "--patch", "patch.json",
                "--output", "assessed.json", "--model", "test"]
        with self.assertRaises(SystemExit) as cm:
            RJ.log_run_step(argv)
        self.assertEqual(cm.exception.code, 1)
        self.assertFalse((run / "assessed.json").exists())
        self.assertFalse((run / "judgement_log").exists())
        changed = dict(entry, assessor_note="Reworded.", questions=[{"question_text": "How is risk owned and reviewed?"}])
        (run / "patch.json").write_text(json.dumps([changed]))
        RJ.log_run_step(argv)
        self.assertEqual(json.loads((run / "assessed.json").read_text())[0]["assessor_note"], "Reworded.")
        self.assertEqual(len(list((run / "judgement_log").glob("*_assessor.json"))), 1)

    def test_a_reviewer_cannot_change_the_facts_it_is_given(self):
        import record_judgements as RJ
        csf = {"subcategory_id": "GV.RM-01", "function": "GV", "outcome": "o", "client": "Sector: Financial.",
               "questions": [{"question_text": "How is risk managed?"}], "question_type": "full_discovery",
               "assessor_note": None}
        self.assertTrue(RJ.validate_patch([csf], [dict(csf, client="Sector: Healthcare.", assessor_note="x")]))
        ai = {"subcategory_id": "MAP 1.1", "bank": "areas", "function": "MAP", "outcome": "o",
              "question_type": "full_discovery", "assessor_note": None,
              "questions": [{"question_text": "How?", "applies_to": "Agents", "rows": ["agent|MAP 1.1"],
                             "criticality_frame": "required"}]}
        reworded = dict(ai, assessor_note="Reworded.", questions=[dict(ai["questions"][0], question_text="How now?")])
        self.assertEqual(RJ.validate_patch([ai], [reworded]), [])
        for change in ({"criticality_frame": "optional"}, {"applies_to": "Model"}, {"rows": []}):
            bad = dict(reworded, questions=[dict(reworded["questions"][0], **change)])
            self.assertTrue(RJ.validate_patch([ai], [bad]), change)
        self.assertTrue(RJ.validate_patch([ai], [dict(reworded, questions=reworded["questions"] * 2)]))

    def test_the_checker_flags_rewrites_it_does_not_confirm(self):
        import record_judgements as RJ
        run = Path(tempfile.mkdtemp())
        mk = lambda sid: {"subcategory_id": sid, "function": "GV", "outcome": "o", "question_type": "full_discovery",  # noqa: E731
                          "questions": [{"question_text": f"How is {sid} done?"}], "evidence_prompt": "",
                          "scope_reason": "", "confidence_flag": "", "assessor_note": None}
        a, b = mk("GV.RM-01"), mk("GV.RM-02")
        (run / "draft.json").write_text(json.dumps([a, b]))
        pa = dict(a, assessor_note="Reworded A.", questions=[{"question_text": "Who owns A and how is it reviewed?"}])
        pb = dict(b, assessor_note="Reworded B.", questions=[{"question_text": "Do you have B?"}])
        (run / "patch.json").write_text(json.dumps([pa, pb]))
        w = lambda name, obj: (run / name).write_text(json.dumps(obj))  # noqa: E731
        argv = lambda *extra: [str(run), "--step", "assessor", "--base", "draft.json", "--patch", "patch.json",  # noqa: E731
                               "--output", "assessed.json", "--model", "test", "--check", "check.json", *extra]
        for bad in ([{"subcategory_id": "GV.RM-01", "verdict": "pass", "reason": ""}],             # one verdict missing
                    [{"subcategory_id": "GV.RM-01", "verdict": "pass", "reason": ""},
                     {"subcategory_id": "GV.RM-02", "verdict": "fail", "reason": ""}]):            # a fail without reason
            w("check.json", bad)
            with self.assertRaises(SystemExit):
                RJ.log_run_step(argv())
        check1 = [{"subcategory_id": "GV.RM-01", "verdict": "pass", "reason": ""},
                  {"subcategory_id": "GV.RM-02", "verdict": "fail", "reason": "Check 2: yes / no."}]
        w("check.json", check1)
        RJ.log_run_step(argv())
        out = {e["subcategory_id"]: e for e in json.loads((run / "assessed.json").read_text())}
        self.assertEqual(out["GV.RM-01"]["confidence_flag"], "")
        self.assertTrue(out["GV.RM-02"]["confidence_flag"].startswith("LOW"))
        self.assertEqual(out["GV.RM-02"]["assessor_note"], "Reworded B. | Checker: Check 2: yes / no.")
        # round 2: an unchanged entry that passed round 1 must still pass
        w("patch1.json", [pa, pb]), w("check1.json", check1)
        w("check.json", [{"subcategory_id": "GV.RM-01", "verdict": "fail", "reason": "Check 3: now chained."},
                         {"subcategory_id": "GV.RM-02", "verdict": "fail", "reason": "Check 2: still yes / no."}])
        with self.assertRaises(SystemExit):
            RJ.log_run_step(argv("--prev-patch", "patch1.json", "--prev-check", "check1.json"))


class Metadata(unittest.TestCase):
    def test_a_refresh_updates_only_the_files_parsed(self):
        import parse_crosswalk as PC
        from datetime import date
        path = Path(tempfile.mkdtemp()) / "metadata.json"
        old = {"last_refreshed": "2020-01-01", "is_placeholder": True, "record_count": None, "notes": "n"}
        path.write_text(json.dumps({"csf2_core.json": dict(old), "other.json": dict(old)}))
        PC.record_refresh(path, {"csf2_core.json": 106})
        doc = json.loads(path.read_text())
        self.assertEqual(doc["csf2_core.json"], {"last_refreshed": date.today().isoformat(), "is_placeholder": False,
                                                  "record_count": 106, "notes": "n"})
        self.assertEqual(doc["other.json"], old)
        self.assertIn("not recorded", PC.record_refresh(path.parent / "missing.json", {"x": 1}))


class Workbook(unittest.TestCase):
    def entries(self, function, sid, text, **extra):
        return [{"subcategory_id": sid, "function": function, "outcome": "o", "question_type": "full_discovery",
                 "questions": [{"question_text": text, "source_reference": "s"}], "evidence_prompt": "ep",
                 "confidence_flag": "", **extra}]

    def write(self, framework, *files, out):
        argv = ["x", *map(str, files), str(out)] + (["--framework", framework] if framework != "csf" else [])
        sys.argv = argv
        WQ.main()

    def test_either_sheet_can_be_rewritten_without_losing_the_other(self):
        import openpyxl
        tmp = Path(tempfile.mkdtemp())
        csf, ai, out = tmp / "csf.json", tmp / "ai.json", tmp / "questionnaire.xlsx"
        csf.write_text(json.dumps(self.entries("GV", "GV.RM-01", "How is risk managed?")))
        ai.write_text(json.dumps(self.entries("GOVERN", "GOVERN 1.1", "How is AI governed?", bank="consumer")))
        self.write("ai", ai, out=out)       # AI first
        self.write("csf", csf, out=out)     # then CSF: used to wipe the AI sheet
        self.write("csf", csf, out=out)     # and a CSF re-run must keep it
        wb = openpyxl.load_workbook(out)
        self.assertEqual(wb.sheetnames, ["CSF Questionnaire", "AI RMF Questionnaire"])
        self.assertEqual(wb["AI RMF Questionnaire"].max_row, 2)
        self.assertEqual(wb["CSF Questionnaire"].max_row, 2)

    def test_the_ai_system_inventory_gets_its_own_table_sheet(self):
        import openpyxl
        from paths import AI_AREA_LABELS, AI_QUESTION_RULES
        tmp = Path(tempfile.mkdtemp())
        ai, out = tmp / "ai.json", tmp / "questionnaire.xlsx"
        ai.write_text(json.dumps(self.entries("GOVERN", "GOVERN 1.1", "How is AI governed?", bank="foundation")))
        self.write("ai", ai, out=out)
        self.assertEqual(openpyxl.load_workbook(out).sheetnames, ["AI RMF Questionnaire"])   # no inventory asked
        inv = json.loads(AI_QUESTION_RULES.read_text())["ai_system_inventory"]
        ai.write_text(json.dumps(self.entries("GOVERN", inv["subcategory"], "In the AI systems sheet, list ...",
                                              bank="foundation")))
        self.write("ai", ai, out=out)
        self.write("ai", ai, out=out)                                    # a re-run replaces it
        wb = openpyxl.load_workbook(out)
        self.assertEqual(wb.sheetnames, ["AI RMF Questionnaire", inv["sheet"]])
        ws = wb[inv["sheet"]]
        heads = [c.value for c in ws[2]]
        labels = list(json.loads(AI_AREA_LABELS.read_text())["labels"].values())
        self.assertEqual(heads[0], inv["system_column"])
        self.assertEqual(sorted(heads[1:1 + len(labels)]), sorted(labels))
        self.assertEqual(heads[1 + len(labels):], list(inv["criticality"]))
        self.assertEqual(len(ws.data_validations.dataValidation), len(heads) - 1)


if __name__ == "__main__":
    unittest.main()
