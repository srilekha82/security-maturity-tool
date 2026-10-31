"""Run: .venv/bin/python -m unittest discover -s tests -v

sheet_io.py: spreadsheet <-> JSON for the sheets in data/common/sheet_specs.json. Everything is written to a temp dir;
the committed ai_areas.xlsx and ai_areas.json are read only."""

import json
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
import sheet_io as S  # noqa: E402
from paths import AI_AREAS, AI_AREAS_XLSX  # noqa: E402


class AiAreasSheet(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.spec = S.load_spec("ai_areas")
        self.xlsx, self.json, self.log = self.tmp / "a.xlsx", self.tmp / "a.json", self.tmp / "audit.jsonl"
        shutil.copy(AI_AREAS_XLSX, self.xlsx)   # the committed spreadsheet is the decision
        self.assertEqual(self.do_import(), 0)

    def edit(self, row, header, value):
        wb = openpyxl.load_workbook(self.xlsx)
        ws = wb[self.spec["sheet"]]
        col = [c.value for c in ws[1]].index(header) + 1
        ws.cell(row=row, column=col, value=value)
        wb.save(self.xlsx)

    def cell(self, row, header):
        ws = openpyxl.load_workbook(self.xlsx)[self.spec["sheet"]]
        return ws.cell(row=row, column=[c.value for c in ws[1]].index(header) + 1).value

    def do_import(self):
        try:
            S.import_sheet(self.spec, self.xlsx, self.json, self.log)
            return 0
        except SystemExit as e:
            return e.code

    def test_committed_json_is_the_import_of_the_committed_spreadsheet(self):
        committed = json.loads(AI_AREAS.read_text())
        self.assertEqual(committed["source_sha"], S.audit.sha(AI_AREAS_XLSX))
        self.assertEqual(json.loads(self.json.read_text())["ai_areas"], committed["ai_areas"])
        ids = [a["id"] for a in committed["ai_areas"]]
        self.assertEqual(len(ids), len(set(ids)))

    def test_spreadsheet_only_columns_stay_out_of_the_json(self):
        first = json.loads(self.json.read_text())["ai_areas"][0]
        self.assertNotIn("example_question", first)
        self.assertEqual(sorted(first), ["applies_to", "covers", "id", "impact", "kind", "name"])

    def test_an_edit_is_stale_until_imported(self):
        self.edit(2, "AI area", "Internal business process")
        self.assertFalse(S.sheet_status("ai_areas", self.xlsx, self.json)[0])
        self.assertEqual(self.do_import(), 0)
        self.assertEqual(json.loads(self.json.read_text())["ai_areas"][0]["name"], "Internal business process")
        self.assertTrue(S.sheet_status("ai_areas", self.xlsx, self.json)[0])

    def test_bad_cells_are_rejected_and_nothing_is_written(self):
        second_id = self.cell(3, "ID")
        for header, value in (("Applies to", "everyone"), ("Kind", "asset"), ("Impact", "data, money"),
                              ("Impact", ""), ("ID", "Bad Id"), ("AI area", ""), ("ID", second_id)):
            self.setUp()
            before = self.json.read_text()
            self.edit(2, header, value)
            self.assertEqual(self.do_import(), 1, (header, value))
            self.assertEqual(self.json.read_text(), before, (header, value))

    def test_impact_is_a_list(self):
        self.edit(2, "Impact", "Access, data, access")
        self.assertEqual(self.do_import(), 0)
        self.assertEqual(json.loads(self.json.read_text())["ai_areas"][0]["impact"], ["access", "data"])

    def test_export_never_overwrites_the_decision(self):
        with self.assertRaises(SystemExit):   # in sync, but the spreadsheet holds a column the JSON does not keep
            S.export_sheet(self.spec, self.xlsx, self.json, self.log)
        self.edit(2, "AI area", "edited")
        with self.assertRaises(SystemExit):   # edited and not imported
            S.export_sheet(self.spec, self.xlsx, self.json, self.log)
        self.assertEqual(self.cell(2, "AI area"), "edited")
        self.assertTrue(self.cell(2, "Example question"))

    def test_wrong_headers_are_rejected(self):
        wb = openpyxl.load_workbook(self.xlsx)
        wb[self.spec["sheet"]].cell(row=1, column=2, value="Surface")
        wb.save(self.xlsx)
        self.assertEqual(self.do_import(), 1)


class MappingSheet(unittest.TestCase):
    """The mapping review sheet: mapping.json -> one line per area x MAP/MEASURE/MANAGE subcategory -> back."""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.spec = S.load_spec("mapping")
        self.xlsx, self.json, self.log = self.tmp / "m.xlsx", self.tmp / "m.json", self.tmp / "audit.jsonl"

        def row(sub, cis=(), status="pass"):
            return {"function": sub.split()[0], "subcategory": sub, "cis_refs": list(cis), "impact": ["access"],
                    "why": f"{sub} checks how the org handles MCP servers and tool credentials.", "status": status,
                    "decided_by": "agent", **({"review_reason": "Too generic."} if status != "pass" else {})}
        rows = [row("MAP 2.2"), dict(row("MAP 4.1", ["MCP-2.1"]), folded=["MAP 1.6"]),
                row("MEASURE 2.7", ["MCP-6.8"], "needs_human_review"),
                row("MEASURE 2.8", ["MCP-8.2"]), row("MANAGE 2.4", ["MCP-6.2"]), row("MANAGE 4.3")]
        self.json.write_text(json.dumps({"areas": {"mcp_tools": {"recorded_on": "2026-10-22", "inputs": {},
                                                                 "rows": rows, "no_question": []}}}))
        self.export()

    def export(self):
        S.export_sheet(self.spec, self.xlsx, self.json, self.log, force=True)

    def do_import(self):
        try:
            S.import_sheet(self.spec, self.xlsx, self.json, self.log)
            return 0
        except SystemExit as e:
            return e.code

    def edit(self, sub, **cells):
        wb = openpyxl.load_workbook(self.xlsx)
        ws = wb[self.spec["sheet"]]
        heads = [c.value for c in ws[1]]
        for r in range(2, ws.max_row + 1):
            if ws.cell(row=r, column=heads.index("Subcategory") + 1).value == sub:
                for header, value in cells.items():
                    ws.cell(row=r, column=heads.index(header) + 1, value=value)
        wb.save(self.xlsx)

    def rows(self):
        return {r["subcategory"]: r for r in json.loads(self.json.read_text())["areas"]["mcp_tools"]["rows"]}

    def test_export_lists_every_mapped_function_subcategory_with_its_text(self):
        wb = openpyxl.load_workbook(self.xlsx)
        lines = list(wb["Mapping"].iter_rows(values_only=True))
        self.assertEqual(len(lines) - 1, 72 - 19 - 12)                  # GOVERN and 12 foundation subs, not in the sheet
        heads = lines[0]
        by_sub = {ln[heads.index("Subcategory")]: ln for ln in lines[1:]}
        self.assertNotIn("GOVERN 1.6", by_sub)
        self.assertNotIn("MANAGE 1.2", by_sub)
        self.assertEqual(by_sub["MAP 2.2"][heads.index("Mapped")], "yes")
        self.assertEqual(by_sub["MEASURE 2.10"][heads.index("Mapped")], "no")
        self.assertTrue(by_sub["MEASURE 2.10"][heads.index("What it means (Playbook outcome)")].startswith("Privacy risk"))
        self.assertIn("MCP-8.2: ", by_sub["MEASURE 2.8"][heads.index("Safeguard titles")])
        cov = list(wb["Coverage"].iter_rows(values_only=True))
        self.assertEqual(len(cov) - 1, 72)                               # every subcategory, GOVERN as foundation
        self.assertEqual(len(cov[0]) - 4, len(json.loads(AI_AREAS.read_text())["ai_areas"]))   # every AI area
        self.assertIn("mcp_tools", cov[0])
        self.assertIn("agent (not drafted)", cov[0])
        cov_by_sub = {r[1]: r for r in cov[1:]}
        self.assertTrue(cov_by_sub["GOVERN 1.6"][3].startswith("foundation"))
        self.assertTrue(cov_by_sub["MANAGE 1.2"][3].startswith("foundation"))
        self.assertEqual(cov_by_sub["MAP 2.2"][3], 1)
        self.assertEqual(cov_by_sub["MEASURE 2.10"][3], 0)

    def test_round_trip_changes_nothing(self):
        before = self.rows()
        self.assertEqual(self.do_import(), 0)
        self.assertEqual(self.rows(), before)

    def test_edits_are_stored_as_human_decisions(self):
        self.edit("MEASURE 2.10", Mapped="yes", Impact="access", **{"CIS safeguards": "MCP-6.8",
                  "Why (comment)": "Personal data a tool can reach is checked for privacy risk."})
        self.edit("MEASURE 2.7", Status="pass")                          # the person decides the failed row
        self.edit("MAP 2.2", Mapped="no")                                # and removes one
        self.assertEqual(self.do_import(), 0)
        rows = self.rows()
        self.assertEqual(rows["MEASURE 2.10"]["decided_by"], "human")
        self.assertEqual(rows["MEASURE 2.10"]["cis_refs"], ["MCP-6.8"])
        self.assertEqual((rows["MEASURE 2.7"]["status"], rows["MEASURE 2.7"]["decided_by"]), ("pass", "human"))
        self.assertNotIn("review_reason", rows["MEASURE 2.7"])
        self.assertNotIn("MAP 2.2", rows)
        self.assertEqual(rows["MANAGE 2.4"]["decided_by"], "agent")      # untouched rows keep what was recorded

    def test_bad_edits_are_rejected_and_nothing_is_written(self):
        cases = [("MEASURE 2.10", {"Mapped": "yes", "Why (comment)": "Privacy of what tools can reach is checked.",
                                   "Impact": "access", "CIS safeguards": "LLM-3.1"}),   # not a candidate
                 ("MEASURE 2.10", {"Mapped": "yes", "Why (comment)": "Privacy of what tools can reach is checked."}),
                 ("MEASURE 2.8", {"Mapped": "no"}),                                     # the audit trail check
                 ("MAP 2.2", {"Subcategory": "GOVERN 1.6"})]                             # foundation, never per area
        for sub, cells in cases:
            self.setUp()
            before = self.json.read_text()
            self.edit(sub, **cells)
            self.assertEqual(self.do_import(), 1, (sub, cells))
            self.assertEqual(self.json.read_text(), before, (sub, cells))

    def test_a_fold_is_a_note_on_both_lines_and_survives_edits(self):
        ws = openpyxl.load_workbook(self.xlsx)["Mapping"]
        heads = [c.value for c in ws[1]]
        notes = {ws.cell(row=r, column=heads.index("Subcategory") + 1).value:
                 ws.cell(row=r, column=heads.index("Why (comment)") + 1).comment for r in range(2, ws.max_row + 1)}
        self.assertIn("Also covers MAP 1.6", notes["MAP 4.1"].text)
        self.assertIn("Folded into MAP 4.1", notes["MAP 1.6"].text)
        self.assertIsNone(notes["MAP 2.2"])
        self.edit("MAP 4.1", **{"Why (comment)": "The MCP servers and tools the AI can call are identified, with what each reaches."})
        self.assertEqual(self.do_import(), 0)
        self.assertEqual((self.rows()["MAP 4.1"]["folded"], self.rows()["MAP 4.1"]["decided_by"]), (["MAP 1.6"], "human"))
        self.export()                                                    # mapping MAP 1.6 on its own ends the fold
        self.edit("MAP 1.6", Mapped="yes", Impact="access", **{"Why (comment)": "The intended access of each tool is documented."})
        self.assertEqual(self.do_import(), 0)
        self.assertNotIn("folded", self.rows()["MAP 4.1"])
        self.assertEqual(self.rows()["MAP 1.6"]["decided_by"], "human")

    def test_removing_a_row_with_folds_is_rejected(self):
        before = self.json.read_text()
        self.edit("MAP 4.1", Mapped="no")                                # MAP 1.6 is folded into it
        self.assertEqual(self.do_import(), 1)
        self.assertEqual(self.json.read_text(), before)

    def test_a_fold_can_be_cleared_in_the_sheet(self):
        self.edit("MAP 4.1", **{"Folded into this row": ""})            # unfold MAP 1.6
        self.assertEqual(self.do_import(), 0)
        self.assertNotIn("folded", self.rows()["MAP 4.1"])
        self.assertEqual(self.rows()["MAP 4.1"]["decided_by"], "human")
        self.export()
        self.edit("MAP 4.1", Mapped="no")                                # now the row can be removed
        self.assertEqual(self.do_import(), 0)
        self.assertNotIn("MAP 4.1", self.rows())

    def test_a_missing_row_is_a_flagged_line_the_person_decides(self):
        doc = json.loads(self.json.read_text())
        doc["areas"]["mcp_tools"]["missing"] = [
            {"function": "MANAGE", "subcategory": "MANAGE 4.1", "reason": "Override and rollback of tool actions."},
            {"function": "MEASURE", "subcategory": "MEASURE 2.6", "reason": "A failed tool fails safely."}]
        doc["areas"]["mcp_tools"]["notes"] = [{"function": "MAP", "subcategory": "MAP 2.2", "reason": "Impact too narrow."}]
        self.json.write_text(json.dumps(doc))
        self.export()
        ws = openpyxl.load_workbook(self.xlsx)["Mapping"]
        heads = [c.value for c in ws[1]]
        line = {ws.cell(row=r, column=heads.index("Subcategory") + 1).value:
                {h: ws.cell(row=r, column=i + 1).value for i, h in enumerate(heads)} for r in range(2, ws.max_row + 1)}
        self.assertEqual((line["MANAGE 4.1"]["Mapped"], line["MANAGE 4.1"]["Status"]), ("no", "needs_human_review"))
        self.assertTrue(line["MANAGE 4.1"]["Evaluator note"].startswith("Missing: "))
        self.assertEqual((line["MAP 2.2"]["Status"], line["MAP 2.2"]["Evaluator note"]), ("pass", "Note: Impact too narrow."))
        self.edit("MANAGE 4.1", Mapped="yes", Impact="access", Status="pass",
                  **{"Why (comment)": "A tool's actions can be overridden and rolled back in operation."})
        self.edit("MEASURE 2.6", Status="pass")                         # decided: not needed
        self.assertEqual(self.do_import(), 0)
        entry = json.loads(self.json.read_text())["areas"]["mcp_tools"]
        self.assertNotIn("missing", entry)
        self.assertEqual(self.rows()["MANAGE 4.1"]["decided_by"], "human")
        self.assertEqual(len(entry["notes"]), 1)

    def test_a_redraft_keeps_what_the_person_decided(self):
        import record_ai_bank as R
        self.edit("MEASURE 2.10", Mapped="yes", Impact="access", **{"Why (comment)": "Personal data a tool can reach is checked."})
        self.edit("MAP 2.2", Mapped="no")
        self.assertEqual(self.do_import(), 0)
        self.assertEqual(json.loads(self.json.read_text())["areas"]["mcp_tools"]["removed_by_human"], ["MAP 2.2"])
        inp = R.load_inputs()
        m = R.mapping_inputs(inp)
        agent = [{"function": r["function"], "subcategory": r["subcategory"], "cis_refs": r["cis_refs"],
                  "impact": r["impact"], "why": "The agent's new wording of what this subcategory checks for tools."}
                 for r in self.rows().values() if r["subcategory"] != "MEASURE 2.10"]
        agent.append(dict(agent[0], subcategory="MAP 2.2", cis_refs=[]))            # the agent brings MAP 2.2 back
        saved = R.AI_MAPPING
        R.AI_MAPPING = self.json
        try:
            _, line = R.store_mapping({"area": "mcp_tools", "rows": agent, "no_question": []}, inp, m, None)
        finally:
            R.AI_MAPPING = saved
        rows = self.rows()
        self.assertNotIn("MAP 2.2", rows)                                             # the person removed it
        self.assertEqual(rows["MEASURE 2.10"]["why"], "Personal data a tool can reach is checked.")
        self.assertEqual(rows["MEASURE 2.10"]["decided_by"], "human")
        self.assertTrue(rows["MANAGE 2.4"]["why"].startswith("The agent's new wording"))
        self.assertIn("2 human decisions kept", line)

    def test_the_mapping_gate(self):
        import ai_pilot as PL
        accepted = self.tmp / "accepted.json"
        saved = PL.AI_MAPPING, PL.AI_MAPPING_ACCEPTED, PL.AI_MAPPING_XLSX
        PL.AI_MAPPING, PL.AI_MAPPING_ACCEPTED, PL.AI_MAPPING_XLSX = self.json, accepted, self.xlsx
        try:
            ok, text = PL.mapping_status()
            self.assertFalse(ok)
            self.assertIn("incomplete (1 of", text)                                    # one area of the 14
            self.assertIn("1 rows need human review", text)
            doc = json.loads(self.json.read_text())
            doc["areas"]["mcp_tools"]["missing"] = [{"function": "MANAGE", "subcategory": "MANAGE 4.1", "reason": "x"}]
            self.json.write_text(json.dumps(doc))
            self.assertIn("1 missing rows to decide", PL.mapping_status()[1])
            del doc["areas"]["mcp_tools"]["missing"]
            self.json.write_text(json.dumps(doc))
            with self.assertRaises(SystemExit):
                PL.accept_mapping(None)
            real = PL.mapping_problems
            PL.mapping_problems = lambda doc: ([], ["MAP 1.2"])                        # as if complete and reviewed
            try:
                self.assertIn("not accepted", PL.mapping_status()[1])
                self.assertIn("reached by no area: MAP 1.2", PL.accept_mapping("looks right"))
                self.assertTrue(PL.mapping_status()[0])
                self.edit("MAP 2.2", Mapped="no")
                self.assertEqual(self.do_import(), 0)
                self.assertIn("changed since accepted", PL.mapping_status()[1])
            finally:
                PL.mapping_problems = real
        finally:
            PL.AI_MAPPING, PL.AI_MAPPING_ACCEPTED, PL.AI_MAPPING_XLSX = saved

    def test_edits_not_imported_block_the_export(self):
        self.edit("MAP 2.2", Mapped="no")
        with self.assertRaises(SystemExit):
            S.export_sheet(self.spec, self.xlsx, self.json, self.log)


class NotesSheet(unittest.TestCase):
    """The impact notes review sheet: notes -> one line per CSF subcategory -> back."""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        long = "Agents can call APIs at machine speed, so permissions sized for humans may be too broad."
        self.doc = {"notes": {
            "PR.AA-05": {"exposure": long, "use": {"base_rag": long, "mcp_agentic": long},
                         "status": "needs_human_review", "review_reason": "use: Check 5: x"},
            "DE.CM-01": {"exposure": long, "status": "pass"}}}
        self.spec = S.load_spec("impact_notes")
        self.xlsx, self.js = self.tmp / "notes.xlsx", self.tmp / "notes.json"
        self.js.write_text(json.dumps(self.doc))
        S.export_sheet(self.spec, self.xlsx, self.js, self.tmp / "audit.jsonl", force=True)

    def edit(self, line, **cells):
        wb = openpyxl.load_workbook(self.xlsx)
        ws = wb[self.spec["sheet"]]
        heads = [c.value for c in ws[1]]
        for header, value in cells.items():
            ws.cell(row=line + 1, column=heads.index(header) + 1, value=value)
        wb.save(self.xlsx)

    def do_import(self):
        try:
            S.import_sheet(self.spec, self.xlsx, self.js, self.tmp / "audit.jsonl")
            return 0
        except SystemExit as e:
            return e.code

    def notes(self):
        return json.loads(self.js.read_text())["notes"]

    def test_round_trip_changes_nothing(self):
        before = self.notes()
        self.assertEqual(self.do_import(), 0)
        self.assertEqual(self.notes(), before)

    def test_an_edit_is_a_human_decision_and_bad_edits_are_rejected(self):
        line = 1 + list(self.doc["notes"]).index("DE.CM-01")
        self.edit(line, **{"Use: model / RAG": "Retrieval must respect the asking user's permissions."})
        self.assertEqual(self.do_import(), 1)                          # one use note without the other
        self.edit(line, **{"Use: tools / agents": "Give each AI tool its own scoped credentials, never a user's."})
        self.assertEqual(self.do_import(), 0)
        self.assertEqual((set(self.notes()["DE.CM-01"]["use"]), self.notes()["DE.CM-01"]["decided_by"]),
                         ({"base_rag", "mcp_agentic"}, "human"))
        pr = 1 + list(self.doc["notes"]).index("PR.AA-05")
        self.edit(pr, Status="pass", **{"Use: model / RAG": "", "Use: tools / agents": ""})
        self.assertEqual(self.do_import(), 0)                          # use layer removed, flag cleared
        self.assertEqual(self.notes()["PR.AA-05"], {"exposure": self.doc["notes"]["PR.AA-05"]["exposure"],
                                                     "status": "pass", "decided_by": "human"})
        self.edit(pr, **{"Exposure (every org)": "too short"})
        self.assertEqual(self.do_import(), 1)

    def test_exposure_is_optional_and_clearing_every_note_drops_the_subcategory(self):
        pr = 1 + list(self.doc["notes"]).index("PR.AA-05")
        self.edit(pr, Status="pass", **{"Exposure (every org)": ""})
        self.assertEqual(self.do_import(), 0)
        self.assertEqual(set(self.notes()["PR.AA-05"]), {"use", "status", "decided_by"})
        S.export_sheet(self.spec, self.xlsx, self.js, self.tmp / "audit.jsonl", force=True)
        de = 1 + list(self.doc["notes"]).index("DE.CM-01")
        self.edit(de, **{"Exposure (every org)": ""})
        self.assertEqual(self.do_import(), 0)
        self.assertNotIn("DE.CM-01", self.notes())


if __name__ == "__main__":
    unittest.main()


class QuestionSheets(unittest.TestCase):
    """The question review sheets: a bank -> one line per question -> back, checked like agent output."""

    def setUp(self):
        import record_ai_bank as R
        from paths import AI_MAPPING
        self.R = R
        real = R.AI_MAPPING
        R.AI_MAPPING = AI_MAPPING                   # the accepted mapping, read only (other tests redirect it)
        self.addCleanup(setattr, R, "AI_MAPPING", real)
        self.tmp = Path(tempfile.mkdtemp())
        self.inp = R.load_inputs()
        ctx = R.question_context(self.inp)
        rows = ctx["rows"]["MEASURE 2.8"]
        groups = {}
        for rid, r in rows.items():
            groups.setdefault(tuple(r["roles"]), []).append(rid)
        q = lambda n, **o: {"question_text": f"For your highest-criticality AI systems, how are AI actions logged ({n}), "  # noqa: E731
                                             "and who reviews the logs?",
                            "playbook_refs": ["MEASURE 2.8 — suggested action 1"], "cis_refs": [],
                            "roles": ["consumer", "producer"], "producer_flags": [], **o}
        areas = {"subcategories": {"MEASURE 2.8": {
            "status": "needs_human_review", "turns": 2, "review_reason": "Q1: vague", **R.rows_record(ctx, "MEASURE 2.8"),
            "questions": [q(n, rows=rids, roles=list(roles)) for n, (roles, rids) in enumerate(groups.items(), 1)]}},
            "summary": {}}
        foundation = {"subcategories": {"GOVERN 1.2": {"status": "pass", "turns": 1, "questions": [
            q("policy", roles=["none", "consumer", "producer"], playbook_refs=["GOVERN 1.2 — suggested action 1"])]}},
            "summary": {}}
        self.files = {}
        for bank, doc in (("areas", areas), ("foundation", foundation)):
            spec = S.load_spec({"areas": "area_questions", "foundation": "foundation_questions"}[bank])
            xlsx, js = self.tmp / f"{bank}.xlsx", self.tmp / f"{bank}.json"
            js.write_text(json.dumps(doc))
            S.export_sheet(spec, xlsx, js, self.tmp / "audit.jsonl", force=True)
            self.files[bank] = (spec, xlsx, js)

    def bank(self, bank):
        return json.loads(self.files[bank][2].read_text())["subcategories"]

    def test_evidence_prompts_are_edited_in_the_sheet_and_a_blank_cell_keeps_them(self):
        spec, xlsx, _ = self.files["foundation"]
        text = "The AI policy, where it is published, and the last record of staff acknowledging it."

        def set_evidence(value):
            wb = openpyxl.load_workbook(xlsx)
            ws = wb[spec["sheet"]]
            heads = [c.value for c in ws[1]]
            ws.cell(row=2, column=heads.index("Evidence prompt") + 1, value=value)
            wb.save(xlsx)
        set_evidence(text)
        self.assertEqual(self.do_import("foundation"), 0)
        self.assertEqual(self.bank("foundation")["GOVERN 1.2"]["questions"][0]["evidence_prompt"], text)
        S.export_sheet(spec, xlsx, self.files["foundation"][2], self.tmp / "audit.jsonl", force=True)
        set_evidence("")
        self.assertEqual(self.do_import("foundation"), 0)
        self.assertEqual(self.bank("foundation")["GOVERN 1.2"]["questions"][0]["evidence_prompt"], text)

    def do_import(self, bank):
        spec, xlsx, js = self.files[bank]
        try:
            S.import_sheet(spec, xlsx, js, self.tmp / "audit.jsonl")
            return 0
        except SystemExit as e:
            return e.code

    def sheet(self, bank):
        spec, xlsx, _ = self.files[bank]
        wb = openpyxl.load_workbook(xlsx)
        ws = wb[spec["sheet"]]
        return wb, ws, [c.value for c in ws[1]]

    def edit(self, bank, line, **cells):
        """Set cells (by header) on data line `line` (1 = first question)."""
        wb, ws, heads = self.sheet(bank)
        for header, value in cells.items():
            ws.cell(row=line + 1, column=heads.index(header) + 1, value=value)
        wb.save(self.files[bank][1])

    def test_export_has_a_line_per_question_and_the_round_trip_changes_nothing(self):
        before = self.bank("areas")
        _, ws, heads = self.sheet("areas")
        self.assertEqual(ws.max_row - 1, len(before["MEASURE 2.8"]["questions"]))
        self.assertIn("Areas covered", heads)
        self.assertEqual(self.do_import("areas"), 0)
        self.assertEqual(self.bank("areas"), before)
        self.assertNotIn("decided_by", self.bank("areas")["MEASURE 2.8"])

    def test_an_edit_is_a_human_decision_and_pass_clears_the_flag(self):
        self.edit("areas", 1, **{"Question": "For your highest-criticality AI systems, how is each AI action traced "
                                             "to the request behind it?"})
        self.assertEqual(self.do_import("areas"), 0)
        e = self.bank("areas")["MEASURE 2.8"]
        self.assertEqual((e["decided_by"], e["status"]), ("human", "needs_human_review"))
        _, ws, heads = self.sheet("areas")
        for line in range(1, ws.max_row):
            self.edit("areas", line, Status="pass")
        self.assertEqual(self.do_import("areas"), 0)
        e = self.bank("areas")["MEASURE 2.8"]
        self.assertEqual(e["status"], "pass")
        self.assertNotIn("review_reason", e)
        self.assertIn("traced", e["questions"][0]["question_text"])

    def test_each_line_shows_only_its_own_evaluator_reason(self):
        js = self.files["areas"][2]
        doc = json.loads(js.read_text())
        doc["subcategories"]["MEASURE 2.8"]["review_reason"] = "Q2: chained; Q0: nothing asks retention; Q3: wrong action"
        js.write_text(json.dumps(doc))
        S.export_sheet(self.files["areas"][0], self.files["areas"][1], js, self.tmp / "audit.jsonl", force=True)
        _, ws, heads = self.sheet("areas")
        cell = lambda line, h: ws.cell(row=line + 1, column=heads.index(h) + 1).value or ""   # noqa: E731
        self.assertEqual((cell(1, "Status"), cell(1, "Evaluator's reason")),
                         ("needs_human_review", "Whole subcategory: nothing asks retention"))
        self.assertEqual(cell(2, "Evaluator's reason"), "chained\nWhole subcategory: nothing asks retention")
        self.assertEqual(self.do_import("areas"), 0)                     # the round trip changes nothing
        self.assertNotIn("decided_by", self.bank("areas")["MEASURE 2.8"])
        for line in range(1, ws.max_row):                                # pass the whole-subcategory failure and Q2
            if line != 3:
                self.edit("areas", line, Status="pass")
        self.assertEqual(self.do_import("areas"), 0)
        e = self.bank("areas")["MEASURE 2.8"]
        self.assertEqual((e["status"], e["review_reason"], e["decided_by"]),
                         ("needs_human_review", "Q0: nothing asks retention; Q3: wrong action", "human"))

    def test_changed_mapping_rows_flag_the_subcategory_until_it_is_set_pass(self):
        js = self.files["areas"][2]
        doc = json.loads(js.read_text())
        e = doc["subcategories"]["MEASURE 2.8"]
        e.update({"status": "pass", "rows_sha": "written-for-old-rows"})
        e.pop("review_reason")
        js.write_text(json.dumps(doc))
        S.export_sheet(self.files["areas"][0], self.files["areas"][1], js, self.tmp / "audit.jsonl", force=True)
        _, ws, heads = self.sheet("areas")
        col = heads.index("Evaluator's reason") + 1
        self.assertTrue(all("mapping rows changed" in ws.cell(row=r, column=col).value for r in range(2, ws.max_row + 1)))
        self.assertEqual(self.do_import("areas"), 0)                     # left flagged: nothing changes
        self.assertEqual(self.bank("areas")["MEASURE 2.8"]["rows_sha"], "written-for-old-rows")
        for line in range(1, ws.max_row):
            self.edit("areas", line, Status="pass")
        self.assertEqual(self.do_import("areas"), 0)                     # a person checked them: stored for the rows now
        e = self.bank("areas")["MEASURE 2.8"]
        ctx = self.R.question_context(self.inp)
        self.assertEqual((e["status"], e["decided_by"], e["rows_sha"]),
                         ("pass", "human", self.R.rows_record(ctx, "MEASURE 2.8")["rows_sha"]))

    def test_a_shared_question_can_be_split_by_adding_a_line(self):
        _, ws, heads = self.sheet("areas")
        line = next(n for n in range(1, ws.max_row)
                    if len(str(ws.cell(row=n + 1, column=heads.index("Areas covered") + 1).value).split(",")) > 1)
        areas = [a.strip() for a in ws.cell(row=line + 1, column=heads.index("Areas covered") + 1).value.split(",")]
        self.edit("areas", line, **{"Areas covered": ", ".join(areas[1:])})
        wb, ws, heads = self.sheet("areas")
        new = [ws.cell(row=line + 1, column=i + 1).value for i in range(len(heads))]
        new[heads.index("Q#")], new[heads.index("Areas covered")] = None, areas[0]
        new[heads.index("Question")] = f"For your highest-criticality AI systems, how is {areas[0]} logged?"
        ws.append(new)
        wb.save(self.files["areas"][1])
        self.assertEqual(self.do_import("areas"), 0)
        qs = self.bank("areas")["MEASURE 2.8"]["questions"]
        self.assertEqual(qs[-1]["rows"], [f"{areas[0]}|MEASURE 2.8"])

    def test_bad_edits_are_rejected_and_nothing_is_written(self):
        _, ws, heads = self.sheet("areas")
        col = lambda h, n: ws.cell(row=n + 1, column=heads.index(h) + 1).value   # noqa: E731
        model_line = next(n for n in range(1, ws.max_row) if "model" in str(col("Areas covered", n)).split(", "))
        cases = {
            "role the area does not allow": (model_line, {"Asked of": "consumer, producer"}),
            "an area dropped, so its row is not asked": (1, {"Areas covered": str(col("Areas covered", 1)).split(", ")[0]
                                                             if "," in str(col("Areas covered", 1)) else "agent, agent"}),
            "an area with no row on the subcategory": (1, {"Areas covered": "code_generation, agent"}),
            "no such suggested action": (1, {"Playbook actions": "MEASURE 2.8 — suggested action 99"}),
            "a subcategory not in the bank": (1, {"Subcategory": "MEASURE 2.7"}),
            "unknown role": (1, {"Asked of": "vendor"}),
        }
        for name, (line, cells) in cases.items():
            before = self.files["areas"][2].read_text()
            self.edit("areas", line, **cells)
            self.assertEqual(self.do_import("areas"), 1, name)
            self.assertEqual(self.files["areas"][2].read_text(), before, name)
            self.setUp()

    def test_deleting_every_line_of_a_subcategory_is_rejected(self):
        wb, ws, _ = self.sheet("foundation")
        ws.delete_rows(2, ws.max_row)
        ws.cell(row=2, column=1, value=None)
        wb.save(self.files["foundation"][1])
        self.assertEqual(self.do_import("foundation"), 1)

    def test_foundation_needs_a_consumer_and_a_producer_question(self):
        self.edit("foundation", 1, **{"Asked of": "none, consumer"})
        self.assertEqual(self.do_import("foundation"), 1)
        self.edit("foundation", 1, **{"Asked of": "none, consumer, producer", "Producer flag": None})
        self.assertEqual(self.do_import("foundation"), 0)

    def test_a_human_decision_is_not_overwritten_by_agent_output(self):
        import argparse
        self.edit("foundation", 1, **{"Question": "Where is your AI policy kept, and how do staff learn of changes?"})
        self.assertEqual(self.do_import("foundation"), 0)
        real = self.R.AI_BANKS
        self.R.AI_BANKS = self.tmp / "banks"
        self.addCleanup(setattr, self.R, "AI_BANKS", real)
        self.R.AI_BANKS.mkdir()
        shutil.copy(self.files["foundation"][2], self.R.AI_BANKS / "question_bank_ai_rmf_foundation.json")
        data = {"GOVERN 1.2": {"status": "pass", "turns": 1, "questions": []}}
        args = argparse.Namespace(bank="foundation", replace_human=False)
        self.assertTrue(self.R.human_decided(data, args))
        self.assertEqual(self.R.human_decided(data, argparse.Namespace(bank="foundation", replace_human=True)), [])
