"""
Usage: .venv/bin/python ./scripts/write_questionnaire.py <questions_json>... <output_xlsx> [--framework csf|ai] [--sheet NAME]

Reads question files (nested format: one entry per subcategory, questions[] array per entry) and writes one sheet of an
XLSX, one row per question; subcategory fields repeat. The last path is the output. One workbook holds the whole
questionnaire, one sheet per framework:

  csf (default)  csf_questions_assessed.json -> questionnaire.xlsx, sheet "CSF Questionnaire".
  ai             ai_questions_assessed_<bank>.json (one or two files: foundation, areas) -> the same questionnaire.xlsx,
                 sheet "AI RMF Questionnaire", with an Area column ("Whole organisation" for a foundation question, the
                 client's AI area, or "This question covers: ..." for a question shared by several areas) and a
                 per-question evidence prompt. When the questions include the AI system inventory
                 (data/ai/question_rules.json ai_system_inventory), it also writes the sheet the client answers it in:
                 one row per AI system, a Yes column per AI area (client labels) and criticality drop-downs.

Either mode opens questionnaire.xlsx if it exists and replaces only its own sheet; the other sheet is never touched, so
the two can be written (or re-written) in any order. The CSF sheet is always first.
"""
import argparse
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).parent))
from options import AI_FUNCTION_ORDER, AI_QUESTION_BANKS, CSF_FUNCTION_ORDER  # noqa: E402
from paths import AI_AREA_LABELS, AI_AREAS, AI_QUESTION_RULES  # noqa: E402
import audit  # noqa: E402

FRAMEWORKS = {
    "csf": {
        "sheet": "CSF Questionnaire",
        "order": CSF_FUNCTION_ORDER,
        "headers": ["Function", "Subcategory ID", "Q#", "Outcome", "Question Type", "Question", "Source", "Confidence",
                    "Evidence Required", "Answer", "Evidence"],
        "widths": [10, 14, 5, 40, 14, 70, 35, 20, 40, 40, 40],
    },
    "ai": {
        "sheet": "AI RMF Questionnaire",
        "order": AI_FUNCTION_ORDER,
        "headers": ["Function", "Subcategory ID", "Area", "Q#", "Outcome", "Question Type", "Question", "Source",
                    "CIS Reference", "Confidence", "Evidence Required", "Answer", "Evidence"],
        "widths": [10, 14, 24, 5, 40, 14, 70, 35, 22, 20, 40, 40, 40],
    },
}


def row_values(framework: str, entry: dict, q_num: int, q: dict) -> list:
    function, sid = entry.get("function", ""), entry.get("subcategory_id", "")
    rest = [entry.get("outcome", ""), entry.get("question_type", ""), q.get("question_text", ""),
            q.get("source_reference", "")]
    confidence, evidence = entry.get("confidence_flag", ""), entry.get("evidence_prompt", "")
    if framework == "ai":
        return ([function, sid, q.get("applies_to", ""), q_num] + rest +
                [", ".join(q.get("cis_refs", [])), confidence, q.get("evidence_prompt") or evidence, "", ""])
    return [function, sid, q_num] + rest + [confidence, evidence, "", ""]


def replace_sheet(wb, name: str):
    if name in wb.sheetnames:  # re-run: replace that sheet, keep the others
        del wb[name]
    return wb.create_sheet(name)


def style_header(cell):
    cell.fill = PatternFill("solid", fgColor="1A56DB")
    cell.font = Font(bold=True, color="FFFFFF")
    cell.alignment = Alignment(wrap_text=True, vertical="top")


def write_inventory(wb, inv: dict) -> None:
    """The AI system inventory as a table: instructions, then a header row, then blank rows with drop-downs."""
    labels = json.loads(AI_AREA_LABELS.read_text())["labels"]
    areas = [labels[a["id"]] for a in json.loads(AI_AREAS.read_text())["ai_areas"]]
    ws = replace_sheet(wb, inv["sheet"])
    headers = [inv["system_column"]] + areas + list(inv["criticality"])
    ws.cell(row=1, column=1, value=inv["instructions"]).font = Font(bold=True)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    first, last = 3, 2 + inv["rows"]
    for col, h in enumerate(headers, 1):
        style_header(ws.cell(row=2, column=col, value=h))
        ws.column_dimensions[get_column_letter(col)].width = 40 if col == 1 else 18
    lists = [["Yes"]] * len(areas) + list(inv["criticality"].values())
    for col, options in enumerate(lists, 2):
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        letter = get_column_letter(col)
        dv.add(f"{letter}{first}:{letter}{last}")
        ws.add_data_validation(dv)
    ws.row_dimensions[2].height = 90
    ws.freeze_panes = "B3"


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("paths", nargs="+", help="question file(s), then the output xlsx")
    parser.add_argument("--framework", choices=list(FRAMEWORKS), default="csf")
    parser.add_argument("--sheet")
    args = parser.parse_args()
    spec = FRAMEWORKS[args.framework]

    if len(args.paths) < 2 or (args.framework == "csf" and len(args.paths) != 2):
        parser.error("give one question file (csf) or one or two (ai), then the output xlsx")
    *question_files, output = args.paths
    questions = [e for f in question_files for e in json.loads(Path(f).read_text())]
    outpath = Path(output)
    ev = audit.begin(audit.run_log(outpath.parent), "write_questionnaire",
                     {"questions": question_files, "output": output, "framework": args.framework})

    sheet_name = args.sheet or spec["sheet"]
    if outpath.is_file():
        wb = openpyxl.load_workbook(outpath)
        ws = replace_sheet(wb, sheet_name)
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name

    for col, h in enumerate(spec["headers"], 1):
        style_header(ws.cell(row=1, column=col, value=h))

    bank_order = {b: i for i, b in enumerate(AI_QUESTION_BANKS)}
    sorted_entries = sorted(
        questions, key=lambda q: (spec["order"].get(q.get("function", ""), 99), q.get("subcategory_id", ""),
                                  bank_order.get(q.get("bank"), 9)))
    row_idx, total_q_rows = 2, 0
    for entry in sorted_entries:
        for q_num, q in enumerate(entry.get("questions", []), 1):
            for col, value in enumerate(row_values(args.framework, entry, q_num, q), 1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            row_idx += 1
            total_q_rows += 1

    for i, w in enumerate(spec["widths"], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    inv = json.loads(AI_QUESTION_RULES.read_text())["ai_system_inventory"] if args.framework == "ai" else None
    if inv and any(e.get("subcategory_id") == inv["subcategory"] for e in questions):
        write_inventory(wb, inv)
    if args.framework == "csf":
        wb.move_sheet(ws, offset=-wb.index(ws))  # CSF sheet first, whichever was written first
    wb.save(outpath)
    print(f"Written: {outpath} [{sheet_name}] ({total_q_rows} question rows across {len(sorted_entries)} subcategories)")
    ev.finish(inputs=question_files, outputs=[outpath],
              detail=f"{sheet_name}: {total_q_rows} question rows, {len(sorted_entries)} subcategories")


if __name__ == "__main__":
    main()
