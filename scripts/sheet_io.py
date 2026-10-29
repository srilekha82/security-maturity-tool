"""
Usage: .venv/bin/python ./scripts/sheet_io.py import <sheet> [--xlsx PATH] [--json PATH]
       .venv/bin/python ./scripts/sheet_io.py export <sheet> [--xlsx PATH] [--json PATH] [--force]

Converts a spreadsheet a person edits into the JSON scripts and agents read, and back. One script for every such sheet:
the columns, required cells and allowed values of each <sheet> live in data/common/sheet_specs.json (today: ai_areas,
mapping, area_questions, foundation_questions). Scripts never read the spreadsheet; they read the JSON.

  ai_areas  the spreadsheet is the decision (data/ai/authoring/ai_areas.xlsx -> data/ai/ai_areas.json).
  mapping   data/ai/mapping.json is the decision; export writes it for review as one line per AI area x MAP, MEASURE or
            MANAGE subcategory (mapped or not, with the Playbook outcome and safeguard titles) plus a Coverage sheet;
            import checks every area with record_ai_bank.py's mapping validator and stores changed lines as
            decided_by human.
  area_questions, foundation_questions
            the question banks (data/ai/banks/question_bank_ai_rmf_<areas|foundation>.json) are the decision; export
            writes one line per question (areas covered, who it is asked of, producer flag, the question, Playbook
            actions, safeguards, status); import checks the whole bank with record_ai_bank.py's questions validator
            and stores changed subcategories as decided_by human. Lines may be edited, added (blank Q#) or deleted;
            subcategories come from generation and cannot be added or dropped here.

  import   XLSX -> JSON. Checks the WHOLE sheet (headers, required cells, allowed values, id format, duplicate ids);
           if anything is wrong it lists the problems, exits 1 and writes nothing. Otherwise writes the JSON with the
           spreadsheet's hash (`source_sha`), which is how the freshness check sees a sheet edited but not imported.
           Columns marked "in_json": false stay in the spreadsheet only.
  export   JSON -> XLSX, with a dropdown on every single-choice column and a "Read me" sheet, then imports the written
           file so the JSON records its hash. Refuses to overwrite a spreadsheet whose edits were never imported (its
           hash differs from the JSON's `source_sha`) unless --force. A spreadsheet that is itself the decision
           (ai_areas) has columns the JSON does not keep, so export refuses to overwrite it at all unless --force.

Column types: id (lowercase letters, digits, _), text, enum (one allowed value), list (one or more allowed values,
comma-separated, lowercased), codes (comma-separated ids kept as written, e.g. MCP-8.2), bool (yes / no).

sheet_status() is what check_freshness.py reports for a sheet.
"""

import argparse
import json
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).parent))
import paths  # noqa: E402
import audit  # noqa: E402

ID_RE = re.compile(r"^[a-z][a-z0-9_]*$")
YES, NO = {"yes", "y", "true"}, {"no", "n", "false"}
README_SHEET = "Read me"


# ── the spec ────────────────────────────────────────────────────────────────

def load_spec(name: str, specs_path: Path | None = None) -> dict:
    specs = json.loads(Path(specs_path or paths.SHEET_SPECS).read_text())["sheets"]
    if name not in specs:
        sys.exit(f"[sheet_io] ERROR: no sheet {name!r} in sheet_specs.json (known: {sorted(specs)})")
    return dict(specs[name], name=name)


def default_paths(spec: dict) -> tuple[Path, Path, Path]:
    """(xlsx, json, audit log) from the constants the spec names in paths.py."""
    return getattr(paths, spec["xlsx"]), getattr(paths, spec["json"]), getattr(paths, spec["audit"]) / "audit.jsonl"


def json_columns(spec: dict) -> list[dict]:
    return [c for c in spec["columns"] if c.get("in_json", True)]


# ── cells ───────────────────────────────────────────────────────────────────

def parse_cell(col: dict, raw) -> tuple[object, list[str]]:
    """(value, problems) for one cell. Blank optional cells become "" (text), [] (list) or None (bool)."""
    text = "" if raw is None else str(raw).strip()
    t = col["type"]
    if not text:
        if col.get("required"):
            return None, ["is required"]
        return ([] if t in ("list", "codes") else None if t == "bool" else ""), []
    if t == "id":
        return text, ([] if ID_RE.match(text) else ["must be lowercase letters, digits and _ (starting with a letter)"])
    if t == "text":
        return text, []
    if t == "enum":
        v = text.lower()
        return v, ([] if v in col["allowed"] else [f"{text!r} is not one of {col['allowed']}"])
    if t == "list":
        vals = []
        for v in (v.strip().lower() for v in text.split(",")):
            if v and v not in vals:
                vals.append(v)
        bad = [v for v in vals if col.get("allowed") and v not in col["allowed"]]
        return vals, ([f"{bad} not in {col['allowed']}"] if bad else [])
    if t == "codes":
        vals = []
        for v in (v.strip() for v in text.split(",")):
            if v and v not in vals:
                vals.append(v)
        return vals, []
    if t == "bool":
        v = text.lower()
        if v in YES or v in NO:
            return v in YES, []
        return None, [f"{text!r} must be yes or no"]
    raise AssertionError(f"unknown column type {t}")


def cell_text(col: dict, value) -> str:
    if col["type"] == "bool":
        return "" if value is None else ("yes" if value else "no")
    if col["type"] in ("list", "codes"):
        return ", ".join(value or [])
    return "" if value is None else str(value)


# ── import ──────────────────────────────────────────────────────────────────

def read_sheet(spec: dict, xlsx: Path) -> tuple[list[dict], list[str]]:
    """(rows, problems) from the spreadsheet; rows are dicts keyed by column key (all columns)."""
    try:
        wb = openpyxl.load_workbook(xlsx, data_only=True)
    except (OSError, ValueError, KeyError) as e:
        return [], [f"cannot open {xlsx}: {e}"]
    if spec["sheet"] not in wb.sheetnames:
        return [], [f"no sheet named {spec['sheet']!r} (sheets: {wb.sheetnames})"]
    rows = list(wb[spec["sheet"]].iter_rows(values_only=True))
    headers = [str(h).strip() if h is not None else "" for h in (rows[0] if rows else [])]
    want = [c["header"] for c in spec["columns"]]
    missing = [h for h in want if h not in headers]
    extra = [h for h in headers if h and h not in want]
    if missing or extra:
        return [], [f"headers must be {want}; missing {missing}, unexpected {extra}"]
    idx = {h: headers.index(h) for h in want}
    out, problems, seen = [], [], {}
    for n, row in enumerate(rows[1:], start=2):
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        rec = {}
        for c in spec["columns"]:
            raw = row[idx[c["header"]]] if idx[c["header"]] < len(row) else None
            value, errs = parse_cell(c, raw)
            problems += [f"row {n}, {c['header']}: {e}" for e in errs]
            rec[c["key"]] = value
        key = rec.get(spec["key"]) if spec["key"] else None   # no single key (mapping): its converter checks repeats
        if key and key in seen:
            problems.append(f"row {n}: {spec['key']} {key!r} already used in row {seen[key]}")
        elif key:
            seen[key] = n
        out.append(rec)
    if not out and not problems:
        problems.append("the sheet has no rows")
    return out, problems


def import_sheet(spec: dict, xlsx: Path, json_path: Path, log: Path) -> str:
    ev = audit.begin(log, f"sheet_io_import_{spec['name']}", {"xlsx": str(xlsx), "json": str(json_path)}, actor="human")
    if not xlsx.is_file():
        ev.fail(f"[sheet_io] ERROR: {xlsx} not found.")
    rows, problems = read_sheet(spec, xlsx)
    if problems:
        ev.reject(problems, xlsx.name, [xlsx], status="rejected")
    if spec.get("shape") == "mapping":
        return import_mapping(spec, xlsx, json_path, ev, rows)
    if spec.get("shape") == "questions":
        return import_questions(spec, xlsx, json_path, ev, rows)
    keep = [c["key"] for c in json_columns(spec)]
    doc = {"description": spec["json_description"], "source": _rel(xlsx), "source_sha": audit.sha(xlsx),
           "generated_by": f"scripts/sheet_io.py import {spec['name']}",
           spec["list_key"]: [{k: r[k] for k in keep} for r in rows]}
    json_path.parent.mkdir(parents=True, exist_ok=True)
    json_path.write_text(json.dumps(doc, indent=2, ensure_ascii=False) + "\n")
    line = f"{len(rows)} {spec['list_key']} from {xlsx.name} -> {_rel(json_path)}"
    print(f"[sheet_io] {line}")
    ev.finish(inputs=[xlsx], outputs=[json_path], detail=line)
    return line


def mapping_context():
    """record_ai_bank.py (the mapping's converter and validator) with the reference data it needs. Imported only for the
    mapping sheet, so ai_areas works without the AI RMF build files."""
    import record_ai_bank as R
    inp = R.load_inputs()
    return R, inp, R.mapping_inputs(inp)


def import_mapping(spec: dict, xlsx: Path, json_path: Path, ev, rows: list[dict]) -> str:
    if not json_path.is_file():
        ev.fail(f"[sheet_io] ERROR: {json_path} not found; draft an area first (/build-ai-bank, Mapping)")
    R, inp, m = mapping_context()
    doc, problems, changed = R.mapping_from_sheet(rows, json.loads(json_path.read_text()), inp, m)
    if problems:
        ev.reject(problems, xlsx.name, [xlsx], status="rejected")
    doc["source_sha"] = audit.sha(xlsx)
    json_path.write_text(json.dumps(doc, indent=2, ensure_ascii=False))
    line = (f"{len(rows)} lines from {xlsx.name}, {changed} changed (decided_by human) -> {_rel(json_path)}; "
            f"{len(R.mapping_coverage(doc, inp, m['rules'])[1])} subcategories not reached yet")
    print(f"[sheet_io] {line}")
    ev.finish(inputs=[xlsx], outputs=[json_path], detail=line)
    return line


def import_questions(spec: dict, xlsx: Path, json_path: Path, ev, rows: list[dict]) -> str:
    if not json_path.is_file():
        ev.fail(f"[sheet_io] ERROR: {json_path} not found; record questions first (/build-ai-bank, Step 5)")
    import record_ai_bank as R
    inp = R.load_inputs()
    doc, problems, changed = R.questions_from_sheet(spec["bank"], rows, json.loads(json_path.read_text()), inp)
    if problems:
        ev.reject(problems, xlsx.name, [xlsx], status="rejected")
    doc["source_sha"] = audit.sha(xlsx)
    json_path.write_text(json.dumps(doc, indent=2, ensure_ascii=False))
    line = (f"{len(rows)} questions from {xlsx.name}, {changed} subcategories changed (decided_by human) -> "
            f"{_rel(json_path)}; {doc['summary']['needs_human_review']} subcategories need human review")
    print(f"[sheet_io] {line}")
    ev.finish(inputs=[xlsx], outputs=[json_path], detail=line)
    return line


# ── export ──────────────────────────────────────────────────────────────────

def write_xlsx(spec: dict, rows: list[dict], xlsx: Path, extra: dict[str, list[list]] | None = None) -> None:
    """Write the spreadsheet: one row per record (missing columns blank), dropdowns, a Read me sheet, and any extra
    read-only sheets ({title: rows of cells, header first})."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = spec["sheet"]
    fill, font = PatternFill("solid", fgColor="1A56DB"), Font(bold=True, color="FFFFFF")
    for i, c in enumerate(spec["columns"], 1):
        cell = ws.cell(row=1, column=i, value=c["header"])
        cell.fill, cell.font, cell.alignment = fill, font, Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(i)].width = c.get("width", 20)
    for r, rec in enumerate(rows, 2):
        for i, c in enumerate(spec["columns"], 1):
            cell = ws.cell(row=r, column=i, value=cell_text(c, rec.get(c["key"])))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if rec.get("_comments", {}).get(c["key"]):   # a note on the cell (hover), not part of the value
                cell.comment = Comment(rec["_comments"][c["key"]], "sheet_io")
    last = max(len(rows) + 1, 2) + 200  # room for new rows
    for i, c in enumerate(spec["columns"], 1):
        choices = c["allowed"] if c["type"] == "enum" else ["yes", "no"] if c["type"] == "bool" else None
        if choices:
            dv = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=not c.get("required"))
            col = get_column_letter(i)
            dv.add(f"{col}2:{col}{last}")
            ws.add_data_validation(dv)
    ws.freeze_panes = spec.get("freeze", "C2")
    ws.auto_filter.ref = f"A1:{get_column_letter(len(spec['columns']))}{max(len(rows) + 1, 2)}"

    for title, cells in (extra or {}).items():
        xs = wb.create_sheet(title)
        for r, values in enumerate(cells, 1):
            for i, v in enumerate(values, 1):
                cell = xs.cell(row=r, column=i, value=v)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if r == 1:
                    cell.fill, cell.font = fill, font
        for i in range(1, len(cells[0]) + 1 if cells else 1):
            xs.column_dimensions[get_column_letter(i)].width = 60 if i == 3 else 14
        xs.freeze_panes = "D2"
        xs.auto_filter.ref = f"A1:{get_column_letter(len(cells[0]) if cells else 1)}{max(len(cells), 2)}"

    rd = wb.create_sheet(README_SHEET)
    for n, line in enumerate(spec.get("readme", []), 1):
        rd.cell(row=n, column=1, value=line)
    start = len(spec.get("readme", [])) + 2
    for i, h in enumerate(("Column", "Type", "Required", "Allowed values", "In JSON", "Help"), 1):
        cell = rd.cell(row=start, column=i, value=h)
        cell.fill, cell.font = fill, font
    for n, c in enumerate(spec["columns"], start + 1):
        allowed = ", ".join(c["allowed"]) if c.get("allowed") else "yes, no" if c["type"] == "bool" else ""
        for i, v in enumerate((c["header"], c["type"], "yes" if c.get("required") else "no", allowed,
                               "yes" if c.get("in_json", True) else "no", c.get("help", "")), 1):
            rd.cell(row=n, column=i, value=v).alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate((22, 10, 10, 40, 9, 80), 1):
        rd.column_dimensions[get_column_letter(i)].width = w
    xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx)


def export_sheet(spec: dict, xlsx: Path, json_path: Path, log: Path, force: bool = False) -> str:
    ev = audit.begin(log, f"sheet_io_export_{spec['name']}", {"xlsx": str(xlsx), "json": str(json_path), "force": force})
    if not json_path.is_file():
        ev.fail(f"[sheet_io] ERROR: {json_path} not found; nothing to export")
    doc = json.loads(json_path.read_text())
    if xlsx.is_file() and not force:
        if audit.sha(xlsx) != doc.get("source_sha"):
            ev.fail(f"[sheet_io] ERROR: {xlsx} has edits that were never imported (its hash differs from the JSON's). "
                    f"Run sheet_io.py import {spec['name']} first, or pass --force to overwrite them.")
        if spec.get("decision", "spreadsheet") == "spreadsheet" and len(json_columns(spec)) < len(spec["columns"]):
            ev.fail(f"[sheet_io] ERROR: {xlsx} is the decision and holds columns the JSON does not keep; exporting "
                    "would blank them. Edit the spreadsheet instead, or pass --force to rebuild it from the JSON.")
    if spec.get("shape") == "mapping":
        R, inp, m = mapping_context()
        rows, coverage = R.mapping_sheet_rows(doc, inp, m)
        write_xlsx(spec, rows, xlsx, {"Coverage": coverage})
    elif spec.get("shape") == "questions":
        import record_ai_bank as R
        rows = R.question_sheet_rows(spec["bank"], doc, R.load_inputs())
        write_xlsx(spec, rows, xlsx)
    else:
        rows = doc[spec["list_key"]]
        write_xlsx(spec, rows, xlsx)
    ev.finish(inputs=[json_path], outputs=[xlsx], detail=f"exported {len(rows)} rows")
    return import_sheet(spec, xlsx, json_path, log)  # records the new file's hash, and proves the round trip


# ── freshness ───────────────────────────────────────────────────────────────

def sheet_status(name: str, xlsx: Path, json_path: Path) -> tuple[bool, str]:
    """(current?, text) for check_freshness.py: the JSON exists and was imported from the spreadsheet as it is now."""
    if not json_path.is_file():
        return False, f"missing (run sheet_io.py import {name})"
    if not xlsx.is_file():
        return False, "spreadsheet missing"
    if json.loads(json_path.read_text()).get("source_sha") != audit.sha(xlsx):
        return False, f"stale (spreadsheet edited since the last import: run sheet_io.py import {name})"
    return True, "current"


def _rel(p: Path) -> str:
    try:
        return str(Path(p).resolve().relative_to(paths.REPO_ROOT))
    except ValueError:
        return str(p)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("mode", choices=["import", "export"])
    ap.add_argument("sheet", help="a sheet name from data/common/sheet_specs.json, e.g. ai_areas")
    ap.add_argument("--xlsx", help="spreadsheet path (default: from the spec)")
    ap.add_argument("--json", help="JSON path (default: from the spec)")
    ap.add_argument("--force", action="store_true", help="export: overwrite the spreadsheet anyway")
    args = ap.parse_args()

    spec = load_spec(args.sheet)
    xlsx, json_path, log = default_paths(spec)
    xlsx, json_path = Path(args.xlsx or xlsx), Path(args.json or json_path)
    if args.mode == "import":
        import_sheet(spec, xlsx, json_path, log)
    else:
        export_sheet(spec, xlsx, json_path, log, args.force)


if __name__ == "__main__":
    main()
