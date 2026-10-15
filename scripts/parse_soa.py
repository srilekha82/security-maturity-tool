"""
Usage: .venv/bin/python ./scripts/parse_soa.py <run_dir> <soa_path>

  run_dir    — the engagement's run folder (created by intake.py)
  soa_path   — path to the SoA Excel file provided by the user

Copies the SoA file into the run folder's
inputs/ subfolder as soa_original.<ext> (client-provided material lives in inputs/),
parses from the copy, and writes soa_parsed.json to the run folder root.
Prints one summary line. This script only extracts; it never assumes. Anything it cannot read
with certainty is left null, and the agent / scope step deals with it.

Output schema:
  [{control_id, control_id_raw, applicable, applicable_raw, exclusion_reason, implemented}]

  control_id       — normalized to match OLIR keys, e.g. "Annex A Controls: 5.1"
  control_id_raw   — exactly as it appears in the SoA file, e.g. "5.1"
  applicable       — True / False, or null if the cell holds a value outside the known yes/no words
  applicable_raw   — the cell as written (so the agent can read an unrecognised value)
  exclusion_reason — string if given, else null
  implemented      — True / False if the SoA has an implementation column with a known value;
                     null (unknown) if there is no such column, the cell is blank or unrecognised,
                     or the control is not applicable

Exits 1, writing nothing, if the layout cannot be read (no ID or no applicability column, or
zero controls parsed). The message lists the headers it saw.
"""

import openpyxl, json, sys, shutil, re
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
import audit  # noqa: E402


ANNEX_A_RANGE = re.compile(r"^([5-9]|[1-9]\d)\.\d")   # 5.x – 99.x → Annex A
MANDATORY_CLAUSE_RANGE = re.compile(r"^([4])\.")        # 4.x → Mandatory Clause

APPLICABLE_YES = {"yes", "y", "applicable", "true", "1", "x"}
APPLICABLE_NO = {"no", "n", "not applicable", "non-applicable", "n/a", "na", "false", "0"}
IMPLEMENTED_YES = {"yes", "y", "implemented", "complete", "completed", "true", "1", "x"}
IMPLEMENTED_NO = {"no", "n", "not implemented", "not started", "false", "0"}


def normalize_control_id(raw: str) -> str:
    """Map SoA control ID to the format used in olir_iso27001.json control_to_csf keys."""
    raw = raw.strip()
    if ANNEX_A_RANGE.match(raw):
        return f"Annex A Controls: {raw}"
    if MANDATORY_CLAUSE_RANGE.match(raw):
        return f"Mandatory Clause: {raw}"
    return raw


def find_controls_sheet(wb):
    """Return the sheet that contains Annex A / SoA controls data."""
    # Strongest signal: sheet name explicitly contains annex or controls
    for name in wb.sheetnames:
        lower = name.lower()
        if "annex" in lower or "controls" in lower:
            return wb[name]
    # Second pass: avoid overview/summary sheets and take first remaining
    skip = ["overview", "revision", "instruction", "cover", "summary"]
    for name in wb.sheetnames:
        if not any(s in name.lower() for s in skip):
            return wb[name]
    return wb.active


def find_header_row(rows, max_scan=10):
    """Return index of the row containing column headers."""
    for i, row in enumerate(rows[:max_scan]):
        norm = [str(c).strip().lower().replace("\n", " ") if c else "" for c in row]
        if any("applicable" in c or "clause" in c or "annex" in c for c in norm):
            return i
    return 0


def col_index(headers, candidates):
    """Return index of first header matching any candidate substring."""
    for cand in candidates:
        for i, h in enumerate(headers):
            if cand in h:
                return i
    return None


def read_flag(cell, yes: set, no: set):
    """True / False for a known yes/no word, None for blank or anything else."""
    if cell is None:
        return None
    s = str(cell).strip().lower()
    if s in yes:
        return True
    if s in no:
        return False
    return None


def main():
    if len(sys.argv) < 3:
        print("Usage: parse_soa.py <run_dir> <soa_path>", file=sys.stderr)
        sys.exit(1)

    run_dir = Path(sys.argv[1])
    soa_path = Path(sys.argv[2])
    ev = audit.begin(audit.run_log(run_dir), "parse_soa", {"soa": str(soa_path)})
    inputs_dir = run_dir / "inputs"
    inputs_dir.mkdir(parents=True, exist_ok=True)

    copied = inputs_dir / f"soa_original{soa_path.suffix}"
    shutil.copy2(soa_path, copied)

    wb = openpyxl.load_workbook(copied, data_only=True)
    ws = find_controls_sheet(wb)

    rows = list(ws.iter_rows(values_only=True))
    header_idx = find_header_row(rows)
    headers = [str(c).strip().lower().replace("\n", " ") if c else "" for c in rows[header_idx]]

    ci = col_index(headers, ["iso/iec 27001", "control id", "control", "clause", "annex"])
    ca = col_index(headers, ["applicable", "applicab", "in scope"])
    # non-applicability reason must come before generic "reason" to avoid matching col 4
    cr = col_index(headers, ["non-applicab", "not applicable", "reason for non", "exclusion", "justification"])
    cs = col_index(headers, ["implemented", "status", "implementation"])

    if ci is None or ca is None:
        ev.fail(f"[parse_soa] ERROR: could not find the {'control ID' if ci is None else 'applicability'} "
                f"column. Sheet '{ws.title}', header row {header_idx}, headers seen: "
                f"{[h for h in headers if h]}. Nothing written.", inputs=[copied])

    result = []
    for row in rows[header_idx + 1:]:
        if not any(row):
            continue

        raw_id = str(row[ci]).strip() if row[ci] else ""
        if not raw_id or raw_id.lower() in ("none", "nan"):
            continue
        # Skip domain/section header rows (no applicable value)
        if row[ca] is None:
            continue

        applicable_raw = str(row[ca]).strip()
        applicable = read_flag(row[ca], APPLICABLE_YES, APPLICABLE_NO)

        reason = str(row[cr]).strip() if cr is not None and row[cr] else ""
        # Ignore boilerplate "refer to column E" filler text
        if reason.lower().startswith("refer to column"):
            reason = ""

        implemented = None
        if applicable and cs is not None:
            implemented = read_flag(row[cs], IMPLEMENTED_YES, IMPLEMENTED_NO)

        result.append({
            "control_id":      normalize_control_id(raw_id),
            "control_id_raw":  raw_id,
            "applicable":      applicable,
            "applicable_raw":  applicable_raw,
            "exclusion_reason": reason or None,
            "implemented":     implemented,
        })

    if not result:
        ev.fail(f"[parse_soa] ERROR: 0 controls parsed. Sheet '{ws.title}', header row {header_idx}, "
                f"headers seen: {[h for h in headers if h]}. Nothing written.", inputs=[copied])

    parsed_path = run_dir / "soa_parsed.json"
    parsed_path.write_text(json.dumps(result, indent=2))

    n = len(result)
    n_app = sum(1 for r in result if r["applicable"] is True)
    n_excl = sum(1 for r in result if r["applicable"] is False)
    n_unk = sum(1 for r in result if r["applicable"] is None)
    n_impl = sum(1 for r in result if r["applicable"] and r["implemented"] is not None)
    impl_col = f"'{headers[cs]}'" if cs is not None else "none"
    summary = (f"[parse_soa] sheet '{ws.title}', id col '{headers[ci]}', applicable col '{headers[ca]}', "
               f"implemented col {impl_col} | {n} controls: {n_app} applicable "
               f"({n_impl} with known implementation), {n_excl} excluded, {n_unk} unrecognised "
               f"applicability → {parsed_path}")
    print(summary)
    ev.finish(inputs=[copied], outputs=[parsed_path], detail=summary)


if __name__ == "__main__":
    main()
